# -*- coding: utf-8 -*-
"""
从 v2 xlsx 生成「智能体页」与「数据总览页」的数据域，供 build_real.py 合并进 data.js。

输出：build/sources/agents_overview.json
  {
    "agentMonthly": [...],        # 智能体页主表（一行=月份×应用）
    "overviewByProvince": {...},  # 数据总览页（区域 -> 月度数组）
    "LY_OVERVIEW_PANORAMA": {...} # 推广/优秀/双周优秀案例总表
  }

数据源（已固化到 build/sources/，不依赖微信临时目录）：
  【智能体】灵运BI重要数据模拟-v2.xlsx  -> Sheet「智能体清单（各省）」A1:AD
  灵运BI数据总览模拟数据_v2.xlsx        -> 「关键数据总览（月）（全网/各省）」+「推广+双周+优秀（总表）」

要点：
  - 总览两个 sheet 是三行前置（第1行分组 / 第2行真实表头 / 第3行"计算口径"说明），
    真实数据从第4行起；这里按"找表头行 + 只保留月份形如 N月 的行"自动跳过说明行。
  - 「是否xx智能体」类表头是 0/1 标记，0=不是、1=是，原样读入并用于统计。
"""
import os
import re
import json

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "sources")

AG_XLSX = os.path.join(SRC, "【智能体】灵运BI重要数据模拟-v2.xlsx")
OV_XLSX = os.path.join(SRC, "灵运BI数据总览模拟数据_v2.xlsx")
OUT_JSON = os.path.join(SRC, "agents_overview.json")

MONTHS6 = ["1月", "2月", "3月", "4月", "5月", "6月"]


# ---------- 小工具 ----------
def num(v):
    """数字：空值返回 None（前端渲染为 /），不填 0。"""
    if v is None or v == "":
        return None
    try:
        f = float(v)
        return int(f) if f == int(f) else round(f, 4)
    except Exception:
        return None


def flag(v):
    """0/1 标记列：0=不是，1=是；空或异常按 0。"""
    try:
        return 1 if int(float(v)) == 1 else 0
    except Exception:
        return 0


def s(v):
    return "" if v is None else str(v).strip()


def month_num(m):
    mm = re.match(r"(\d+)", str(m))
    return int(mm.group(1)) if mm else 999


# ---------- 1) 智能体月度宽表 ----------
def build_agent_monthly():
    """
    Sheet「智能体清单（各省）」列位（0 基）：
      0月份 1省份 2创建人 3创建时间 4应用名称 5应用说明 6应用类型 7应用标签 8应用状态
      9应用调用量
      16是否有调用标记 17是否活跃 18是否活跃且高提效 19是否高热度
      20是否推广 21是否优秀 22是否双周优秀
      23提效比率 24节约人年 25复制量 26是否使用大模型 27应用成功率 28平均响应耗时 29单笔节约时长
    """
    wb = openpyxl.load_workbook(AG_XLSX, read_only=True, data_only=True)
    ws = wb["智能体清单（各省）"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or r[0] is None:
            continue
        mraw = s(r[0])
        if not re.match(r"^\d+月$", mraw):
            continue
        calls_val = num(r[9]) or 0
        # T列「是否高热度」存在 Excel 公式缓存缺失风险（交接手册 §6.4），
        # 该列为空时按口径 调用量>100000 重新计算，避免整列为 0。
        is_hot = flag(r[19]) if r[19] is not None else (1 if calls_val > 100000 else 0)
        out.append({
            "month": mraw,
            "province": s(r[1]) or "未标注",
            "creator": s(r[2]),
            "createdAt": s(r[3]),
            "name": s(r[4]),
            "description": s(r[5]),
            "type": s(r[6]),
            "tags": s(r[7]),
            "status": s(r[8]),
            "calls": calls_val,
            # 「是否xx」0/1 标记列
            "isCalled": flag(r[16]),
            "isActive": flag(r[17]),
            "isActiveHighEff": flag(r[18]),
            "isHot": is_hot,
            "isPromo": flag(r[20]),
            "isExcellent": flag(r[21]),
            "isBiweek": flag(r[22]),
            "isLLM": flag(r[26]),
            # 数值列
            # 注：节约人年(Y/24) 与 单笔节约时长(AD/29) 按交接手册 §5.1「不必进入前端精简对象，
            # 以控制 data.js 大小」不写入；若后续页面需要，在此恢复即可。
            "effRatio": num(r[23]),
            "copy": num(r[25]) or 0,
            "success": num(r[27]),
            "response": num(r[28]),
        })
    wb.close()
    return out


# ---------- 2) 数据总览（全网 / 各省）----------
def build_overview_by_province(wb):
    """
    列位（0 基，取自第2行真实表头）：
      0月份 1省份 2智能体调用量 3总Token数 4忙时Token 5闲时Token 6累计注册人数 7平台总访问PV
      8投产率 9活跃智能体数 10高热度智能体数 11推广智能体数 12优秀智能体数 13双周优秀智能体数
      14智能体总数量（上线） 15新增注册人数 16活跃用户覆盖度（月） 17WAU周活用户
      18活跃用户数（月） 19沉默/流失用户数 20节约人年 21模型计费 22忙时模型计费 23闲时模型计费
    """
    merged = {}
    seen = set()
    # 「全网」以「关键数据总览（月）（全网）」为准；各省表里也含「全网」行，需跳过避免重复累加。
    for sheet_name, allow_national in (
        ("关键数据总览（月）（全网）", True),
        ("关键数据总览（月）（各省）", False),
    ):
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        # 定位真实表头行（cell0 == "月份"）
        hdr = None
        for i, r in enumerate(rows[:5]):
            if r and s(r[0]) == "月份":
                hdr = i
                break
        if hdr is None:
            print(f"  [overview] {sheet_name}: 未找到表头行，跳过")
            continue

        for r in rows[hdr + 1:]:
            if r is None or r[0] is None:
                continue
            mraw = s(r[0])
            # 跳过「计算口径」说明行等非数据行
            if not re.match(r"^\d+月$", mraw):
                continue
            prv = r[1]
            if prv in (None, "", "总计", "全网"):
                prv = "全网"
            else:
                prv = s(prv)
            if prv == "全网" and not allow_national:
                continue  # 各省表的「全网」行不参与，避免与全网表重复
            key = (prv, mraw)
            if key in seen:
                continue  # 同区域同月去重兜底
            seen.add(key)
            rec = {
                "month": mraw,
                "calls": num(r[2]),
                "tokens": num(r[3]),
                "tokensBusy": num(r[4]),
                "tokensIdle": num(r[5]),
                "cumRegUsers": num(r[6]),
                "pagePV": num(r[7]),
                "productionRate": num(r[8]),
                "activeAgents": num(r[9]),
                "hotAgents": num(r[10]),
                "promoAgents": num(r[11]),
                "excellentAgents": num(r[12]),
                "excellentAgentsBiweek": num(r[13]),
                "totalAgentsOnline": num(r[14]),
                "newRegUsers": num(r[15]),
                "activeCover": num(r[16]),
                "activeUsersWAU": num(r[17]),
                "activeUsers": num(r[18]),
                "silentLost": num(r[19]),
                "personYear": num(r[20]),
                "cost": num(r[21]),
                "costBusy": num(r[22]),
                "costIdle": num(r[23]),
                # 文档未提供的派生项，保留字段以兼容前端
                "producingAgents": None,
                "activeUsersCum": None,
                "activeCoverCum": None,
            }
            merged.setdefault(prv, []).append(rec)

    for k in merged:
        merged[k].sort(key=lambda x: month_num(x["month"]))
    return merged


# ---------- 3) 推广/优秀/双周优秀 案例总表 ----------
def build_panorama(wb):
    ws = wb["推广+双周+优秀（总表）"]
    rows = list(ws.iter_rows(values_only=True))
    prows = []
    scenes = set()
    for r in rows[1:]:
        if r is None or r[0] is None:
            continue
        tags = s(r[0])
        scene = s(r[1])
        if scene:
            scenes.add(scene)
        rec = {
            "tags": tags,
            "scene": scene,
            "province": s(r[2]),
            "name": s(r[3]),
            "createdAt": s(r[4]),
            "creator": s(r[5]),
        }
        for i, m in enumerate(MONTHS6):
            rec[m] = num(r[6 + i])
        prows.append(rec)
    return {"rows": prows, "months": MONTHS6, "scenes": sorted(scenes)}


def main():
    agent_monthly = build_agent_monthly()

    wb = openpyxl.load_workbook(OV_XLSX, read_only=True, data_only=True)
    overview_by_province = build_overview_by_province(wb)
    panorama = build_panorama(wb)
    wb.close()

    payload = {
        "agentMonthly": agent_monthly,
        "overviewByProvince": overview_by_province,
        "LY_OVERVIEW_PANORAMA": panorama,
    }
    os.makedirs(SRC, exist_ok=True)
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)

    # ---- 验收打印 ----
    print(f"[gen_agents_overview] agentMonthly = {len(agent_monthly)} 条")
    by_m = {}
    for r in agent_monthly:
        b = by_m.setdefault(r["month"], {"n": 0, "calls": 0, "act": 0, "hot": 0})
        b["n"] += 1
        b["calls"] += r["calls"] or 0
        b["act"] += r["isActive"]
        b["hot"] += r["isHot"]
    print("  月 | 行数 | 调用量 | 活跃(0/1标记) | 高热度(0/1标记)")
    for m in sorted(by_m, key=month_num):
        b = by_m[m]
        print(f"  {m} | {b['n']} | {b['calls']} | {b['act']} | {b['hot']}")

    print(f"[gen_agents_overview] overviewByProvince = {len(overview_by_province)} 个区域")
    net = overview_by_province.get("全网", [])
    print(f"  全网月份: {[x['month'] for x in net]}")
    if net:
        r0 = net[0]
        print("  全网首月关键字段:",
              {k: r0.get(k) for k in ("calls", "productionRate", "activeAgents", "hotAgents",
                                      "promoAgents", "excellentAgents", "excellentAgentsBiweek",
                                      "newRegUsers", "activeCover", "activeUsers", "personYear")})
    print(f"[gen_agents_overview] PANORAMA rows={len(panorama['rows'])} scenes={len(panorama['scenes'])}")
    print("已写出:", OUT_JSON)


if __name__ == "__main__":
    main()
