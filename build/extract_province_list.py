# -*- coding: utf-8 -*-
"""
从《【合】灵运BI重要数据（终版）(1).xlsx》的「智能体清单（各省）」sheet 抽取
新增字段：节约金额 / 产生真实价值 / 是否正向价值。

仅输出「匹配键 + 3 个新字段」，供 merge_province_list.js 按行匹配合并进现有
agentMonthly 域；不改动调用量、Token、promoScene 等既有字段，以保证已验收的
KPI 与案例联动不被触动。

匹配键：月份 + 应用ID + 省份 + 应用名称 + 创建人（归一化，去除空白）。
"""
import json
import os
import re

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = "/Users/mac/Downloads/【合】灵运BI重要数据（终版） (1).xlsx"
OUT = os.path.join(BASE, "build", "province_list_newfields.json")

REPORT_MONTHS = ["1月", "2月", "3月", "4月", "5月", "6月", "7月", "8月"]
MONTH_RE = re.compile(r"^(\d{1,2})月$")


def num(x, nd=4):
    if x is None:
        return None
    if isinstance(x, bool):
        return int(x)
    if isinstance(x, (int, float)):
        if x != x:
            return None
        return round(float(x), nd)
    s = str(x).strip()
    if s in ("", "/", "暂无数据", "None", "nan", "-"):
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return round(f, nd)


def i01(x):
    return 1 if num(x) else 0


def txt(x):
    if x is None:
        return ""
    if isinstance(x, str):
        s = x.strip()
        return "" if s in ("None", "nan") else s
    return str(x).strip()


def norm_key(s):
    return re.sub(r"\s+", "", txt(s))


def month_num(m):
    mm = MONTH_RE.match(txt(m))
    return int(mm.group(1)) if mm else 99


def main():
    import openpyxl
    if not os.path.exists(SRC):
        raise SystemExit("数据源不存在：" + SRC)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    if "智能体清单（各省）" not in wb.sheetnames:
        raise SystemExit("找不到工作表：智能体清单（各省）")
    ws = wb["智能体清单（各省）"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    H = {str(v).strip(): i for i, v in enumerate(header) if v is not None and str(v).strip()}
    REQUIRED = ["月份", "省份", "应用名称", "应用ID", "创建人", "节约金额", "产生真实价值", "是否正向价值"]
    miss = [k for k in REQUIRED if k not in H]
    if miss:
        raise SystemExit("缺少必需表头，已中止：" + "、".join(miss))

    def g(row, key):
        i = H.get(key)
        return row[i] if i is not None else None

    out_rows = []
    month_seen = set()
    pos_yes = 0
    for r in rows[1:]:
        month = txt(g(r, "月份"))
        name = txt(g(r, "应用名称"))
        if not month or not name:
            continue
        month_seen.add(month)
        province = txt(g(r, "省份")) or "未标注"
        creator = txt(g(r, "创建人"))
        app_id = txt(g(r, "应用ID")) or ""
        save_amount = num(g(r, "节约金额"), 2)
        real_value = num(g(r, "产生真实价值"), 2)
        positive_value = i01(g(r, "是否正向价值"))
        if positive_value:
            pos_yes += 1
        key = "|".join([
            month, norm_key(app_id), norm_key(province), norm_key(name), norm_key(creator)
        ])
        out_rows.append({
            "key": key,
            "saveAmount": save_amount,
            "realValue": real_value,
            "positiveValue": positive_value,
        })

    data = {
        "reportMonths": REPORT_MONTHS,
        "rows": out_rows,
        "stats": {
            "total": len(out_rows),
            "positiveYes": pos_yes,
            "positiveNo": len(out_rows) - pos_yes,
            "months": sorted(month_seen, key=month_num),
        },
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print("抽取行数：%d" % len(out_rows))
    print("月份分布：%s" % data["stats"]["months"])
    print("是否正向价值 是：%d ｜ 否：%d" % (pos_yes, len(out_rows) - pos_yes))
    print("输出：%s（%.2f MB）" % (OUT, os.path.getsize(OUT) / 1048576))


if __name__ == "__main__":
    main()
