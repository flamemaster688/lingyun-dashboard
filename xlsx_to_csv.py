# -*- coding: utf-8 -*-
"""
把金山文档导出的 xlsx 整文件解析为 build_real.py 需要的 _raw_sheetN.csv。
相比 get_range_data，整文件下载不受「单表行数上限」限制，能拿到完整 178 行智能体表。

用法：
  python xlsx_to_csv.py <file.xlsx> [输出目录]
按工作表名映射到 _raw_sheetN.csv（原始网格，含表头，值与 get_range_data 取值一致）。
"""
import sys, csv, os

NAME_MAP = {
    "灵运平台等效人年(推广+优秀+双周)": "_raw_sheet4.csv",
    "1-7月调用量+节约人年": "_raw_sheet5.csv",
    "各省调用量+活跃": "_raw_sheet6.csv",
    "各省Token+费用": "_raw_sheet7.csv",
    "多能力等效人年汇总": "_raw_sheet8.csv",
    "0720-0726埋点日志": "_raw_sheet9.csv",
}


def sheet_to_grid(ws):
    """遍历单元格，按绝对行列位置重建网格（空单元格留空）。"""
    cells = {}
    max_r = -1
    max_c = -1
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            r = cell.row - 1
            c = cell.column - 1
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            elif isinstance(v, float):
                v = repr(v)
            cells[(r, c)] = "" if v is None else str(v)
            max_r = max(max_r, r)
            max_c = max(max_c, c)
    if max_r < 0 or max_c < 0:
        return [[""]]
    grid = [["" for _ in range(max_c + 1)] for _ in range(max_r + 1)]
    for (r, c), v in cells.items():
        grid[r][c] = v
    return grid


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("usage: xlsx_to_csv.py <file.xlsx> [outdir]\n")
        sys.exit(2)
    xlsx = sys.argv[1]
    outdir = sys.argv[2] if len(sys.argv) > 2 else "."
    os.makedirs(outdir, exist_ok=True)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    sheets = {ws.title: ws for ws in wb.worksheets}
    for name, out in NAME_MAP.items():
        ws = sheets.get(name)
        if ws is None:
            for t, w in sheets.items():
                if name[:4] in t or t[:4] in name:
                    ws = w
                    break
        if ws is None:
            print(f"WARN: 未找到工作表「{name}」，跳过")
            continue
        grid = sheet_to_grid(ws)
        path = os.path.join(outdir, out)
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            for row in grid:
                w.writerow(row)
        print(f"OK: 「{name}」 -> {out} ({len(grid)} 行 × {len(grid[0]) if grid else 0} 列)")
    wb.close()


if __name__ == "__main__":
    main()
